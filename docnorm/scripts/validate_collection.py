"""Validate the delivered export run, including a second deterministic run.

Run from any directory using the docnorm virtual environment.
"""
import hashlib
import json
from pathlib import Path
import re
from urllib.parse import unquote, urlsplit
import xml.etree.ElementTree as ET
import zipfile

from markdown_it import MarkdownIt

from docnorm.engine import normalize


def hashes(root):
    return {p.relative_to(root).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(root.rglob("*")) if p.is_file() and not p.is_symlink()}


def links(text):
    for token in MarkdownIt().parse(text):
        for child in token.children or []:
            if child.type == "link_open":
                yield child.attrGet("href")
            elif child.type == "image":
                yield child.attrGet("src")


def main():
    package = Path(__file__).resolve().parents[1]
    source = package.parent / "drive-download-20261012T111012Z-1-001"
    output = package.parent / "normalized"
    original = json.loads((package / "input-hashes.json").read_text())
    assert hashes(source) == original, "input files changed before validation"
    before = hashes(output)
    report = normalize(source, output)
    assert hashes(output) == before, "rerun output is not byte-identical"
    assert hashes(source) == original, "input files changed during validation"
    manifest = [json.loads(line) for line in (output / "manifest.jsonl").read_text().splitlines()]
    assert {r["path"] for r in manifest} == set(original), "manifest does not cover inputs"
    assert all(r["status"] != "pending" for r in manifest)
    office_text_runs = 0
    worksheets = 0
    def plain(text):
        return " ".join(re.sub(r"[\\*_`]", "", text).split())
    for record in manifest:
        if record["path"].endswith(".docx"):
            text = plain((output / record["output"]).read_text())
            with zipfile.ZipFile(source / record["path"]) as archive:
                xml = ET.fromstring(archive.read("word/document.xml"))
            for node in xml.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
                if node.text and node.text.strip():
                    assert plain(node.text) in text, f"missing DOCX text: {record['path']}: {node.text}"
                    office_text_runs += 1
        elif record["path"].endswith(".xlsx"):
            from openpyxl import load_workbook
            book = load_workbook(source / record["path"], read_only=True, keep_links=False)
            try:
                provenance = json.loads((output / Path(record["output"]).parent / "provenance.json").read_text())
                assert book.sheetnames == [s["name"] for s in provenance["locators"]["worksheets"]]
                worksheets += len(book.sheetnames)
            finally:
                book.close()
    roots = json.loads((output / "roots.json").read_text())["roots"]
    provenance_count = 0
    for path in sorted((output / "documents").rglob("provenance.json")):
        p = json.loads(path.read_text())
        assert (output / p["normalized_path"]).is_file()
        assert path.parent.name == p["document_id"]
        for ref in p["sources"]:
            original_path = Path(roots[ref["root_id"]]["path"]) / ref["path"]
            assert original_path.is_file()
            assert hashlib.sha256(original_path.read_bytes()).hexdigest() == ref["sha256"]
        provenance_count += 1
    assets = json.loads((output / "assets/provenance.json").read_text())["assets"]
    for asset in assets:
        assert hashlib.sha256((output / asset["path"]).read_bytes()).hexdigest() == asset["sha256"]
        for ref in asset["sources"]:
            assert original[ref["path"]] == ref["sha256"]
    catalog_links = list(links((output / "llms.txt").read_text()))
    for target in catalog_links:
        assert (output / unquote(urlsplit(target).path)).is_file(), target
    generated_links = 0
    for path in (output / "documents").rglob("content.md"):
        for target in links(path.read_text()):
            parsed = urlsplit(target)
            if parsed.scheme or not parsed.path:
                continue
            local = path.parent / unquote(parsed.path)
            if re.fullmatch(r"\.\./[0-9a-f]{24}/content\.md|\.\./\.\./assets/[0-9a-f]{64}\.[a-z]+", parsed.path):
                assert local.is_file(), f"broken rewritten link: {path}: {target}"
                generated_links += 1
    validation = {"source_files_unchanged": True, "source_file_count": len(original),
                  "manifest_complete": True, "provenance_documents_validated": provenance_count,
                  "asset_provenance_validated": len(assets), "catalog_links_validated": len(catalog_links),
                  "rewritten_links_validated": generated_links, "rerun_byte_identical": True,
                  "docx_text_runs_retained": office_text_runs, "xlsx_worksheets_retained": worksheets,
                  "document_ids_stable": True, "normalization_report": str(output / "report.json"),
                  "file_status_counts": report["file_status_counts"],
                  "partial_documents": report["partial_documents"]}
    (package / "validation-report.json").write_text(json.dumps(validation, ensure_ascii=False, sort_keys=True, indent=2) + "\n")
    print(json.dumps(validation, indent=2))


if __name__ == "__main__":
    main()
