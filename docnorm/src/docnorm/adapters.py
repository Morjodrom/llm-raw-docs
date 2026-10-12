"""Format adapters. Only the local MarkItDown HTML/DOCX converters are loaded."""
import csv
import html
import io
import json
import re
import zipfile
from pathlib import Path

from bs4 import BeautifulSoup

from .models import Converted, Fragment
from .normalization import decode, expand_markdown, fence, title_for
from .safety import safe_path


def cell(value):
    return str(value if value is not None else "").replace("\\", "\\\\").replace("|", "\\|").replace("\r\n", "<br>").replace("\n", "<br>").replace("\r", "<br>")


def table(rows, header=True):
    rows = [list(row) for row in rows]
    if not rows:
        return ""
    width = max(map(len, rows))
    if not width:
        return ""
    if not header:
        rows.insert(0, [f"Column {i + 1}" for i in range(width)])
    padded = [row + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(cell(v) for v in row) + " |" for row in padded]
    lines.insert(1, "| " + " | ".join(["---"] * width) + " |")
    return "\n".join(lines) + "\n"


def local_markitdown(data: bytes, extension: str):
    # Directly use MarkItDown's supported converters: no Magika model, ONNX
    # telemetry, plugin discovery, URL handling or executable/image converters.
    import onnxruntime
    onnxruntime.disable_telemetry_events()
    from markitdown import StreamInfo
    from markitdown.converters import DocxConverter, HtmlConverter
    converter = HtmlConverter() if extension == ".html" else DocxConverter()
    return converter.convert(io.BytesIO(data), StreamInfo(extension=extension, charset="utf-8"),
                             heading_style="ATX", keep_data_uris=True).markdown


def html_to_markdown(text):
    soup = BeautifulSoup(text, "html.parser")
    for node in soup.select("script, style, nav, noscript, .yfm-anchor, .yfm-clipboard-button, [data-no-index='true']"):
        node.decompose()
    anchors = []
    tokens = {}
    for i, node in enumerate(soup.select("[id], a[name]")):
        anchor = node.get("id") or node.get("name")
        if not anchor or anchor in anchors:
            continue
        anchors.append(anchor)
        if anchor.startswith("inline-code-id-"):
            continue  # Diplodoc clipboard-control ID, not document structure.
        token = f"DOCNORMANCHOR{i:06d}TOKEN"
        tokens[token] = f'<a id="{html.escape(anchor, quote=True)}"></a>'
        if node.name in {"td", "th", "li"}:
            node.insert(0, soup.new_string(token))
        elif node.name in {"span", "code", "a", "em", "strong", "b", "i", "img"}:
            node.insert_before(soup.new_string(token))
        else:
            marker = soup.new_tag("p")
            marker.string = token
            node.insert_before(marker)
    for node in soup.select(".yfm-note, .note, [role='alert']"):
        label = node.get("data-note-type") or " ".join(node.get("class", ["Note"]))
        marker = soup.new_tag("p")
        marker.string = f"Note / warning: {label}"
        node.insert(0, marker)
    # Keep collapsed content, exposing the summary as a normal paragraph.
    for node in soup.select("summary"):
        node.name = "p"
    markdown = local_markitdown(str(soup).encode(), ".html")
    for token, anchor in tokens.items():
        markdown = markdown.replace(token, anchor)
    return markdown, anchors


def html_adapter(root, path, warnings):
    text, ws = decode(path.read_bytes())
    warnings.extend(ws)
    soup = BeautifulSoup(text, "html.parser")
    state_node = soup.find("script", id="diplodoc-state")
    metadata, locators = {}, {}
    title = soup.title.get_text() if soup.title else path.stem
    base = path.parent
    if state_node:
        state = json.loads(state_node.string or state_node.get_text())
        data = state["data"]
        metadata = {"meta": data.get("meta", {}), "router": state.get("router", {}), "lang": state.get("lang")}
        title = data.get("title") or title
        if data.get("html", "").strip():
            body, anchors = html_to_markdown(html.unescape(data["html"]))
        elif data.get("leading") and isinstance(data.get("data"), dict):
            leading = data["data"]
            title = leading.get("title") or title
            descriptions = leading.get("description", [])
            if isinstance(descriptions, str):
                descriptions = [descriptions]
            parts = [*descriptions]
            for link in leading.get("links", []):
                parts.append(f"- [{link.get('title', link.get('href', 'Link'))}]({link.get('href', '')})")
                description = link.get("description")
                if description:
                    parts.append("  " + str(description))
            remainder = {k: v for k, v in leading.items() if k not in {"title", "description", "links"}}
            if remainder:
                parts.append(fence(json.dumps(remainder, ensure_ascii=False, indent=2), "json"))
            body, anchors = "\n\n".join(parts), []
        else:
            raise ValueError("Diplodoc article is empty or unsupported")
        locators["html_anchors"] = anchors
        converter = "diplodoc+MarkItDown.HtmlConverter"
    else:
        selected = soup.find("article") or soup.find("main") or soup.body or soup
        body, anchors = html_to_markdown(str(selected))
        locators["html_anchors"] = anchors
        converter = "MarkItDown.HtmlConverter"
    base_tag = soup.find("base", href=True)
    if base_tag:
        href = base_tag["href"]
        if re.match(r"^[a-zA-Z][\w+.-]*:", href) or href.startswith("//"):
            # Materialize external-base links as external URLs without fetching.
            from urllib.parse import urljoin
            body = re.sub(r"(?<=\]\()([^\s)]+)", lambda m: urljoin(href, m[0]), body)
        else:
            try:
                base = safe_path(root, path.parent / href)
            except ValueError:
                warnings.append("HTML base escapes input root; ignored")
    return title, [Fragment(body, base)], converter, metadata, locators


def architecture_adapter(obj, path):
    def readable(value, depth=2):
        if isinstance(value, list):
            return "\n\n".join(readable(item, depth) for item in value)
        if not isinstance(value, dict):
            return str(value)
        parts = []
        label = value.get("title") or value.get("name") or value.get("id")
        if label:
            parts.append("#" * min(depth, 6) + " " + str(label))
        for key, item in value.items():
            if key in {"title", "name"}:
                continue
            if key == "markdown" and item:
                parts.append(item.replace("\\n", "\n") if "\n" not in item else item)
            elif key == "html" and value.get("markdown"):
                continue  # Equivalent rendered representation, not additional data.
            elif key == "html" and item:
                parts.append(html_to_markdown(item)[0])
            elif isinstance(item, (dict, list)):
                parts.append("#" * min(depth + 1, 6) + " " + key)
                parts.append(readable(item, depth + 2))
            else:
                parts.append(f"- **{key}**: {item}")
        return "\n\n".join(parts)

    parts = []
    for key, value in obj.items():
        if key == "title":
            continue
        parts.append(f"## {key}\n\n{readable(value, 3)}")
    return Fragment("\n\n".join(parts), path.parent)


def check_office_archive(path):
    with zipfile.ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist()) > 256 * 1024 * 1024:
            raise ValueError("Office archive exceeds 256 MiB unpacked limit")


def xlsx_adapter(path, warnings):
    from openpyxl import load_workbook
    check_office_archive(path)
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    parts, locators = [], []
    try:
        for sheet in formulas:
            rows = []
            cached_rows = cached[sheet.title].iter_rows()
            for row, values in zip(sheet.iter_rows(), cached_rows):
                cells = []
                for c, v in zip(row, values):
                    if c.data_type == "f":
                        cells.append(v.value if v.value is not None else c.value)
                        if v.value is None:
                            warnings.append(f"worksheet {sheet.title}, {c.coordinate}: formula has no cached value; formula retained")
                    else:
                        cells.append(c.value)
                rows.append(cells)
            while rows and all(v is None for v in rows[-1]):
                rows.pop()
            width = max((i + 1 for row in rows for i, v in enumerate(row) if v is not None), default=0)
            rows = [row[:width] for row in rows]
            parts.append(f"## {sheet.title}\n\n" + (table(rows, header=False) if width else "Empty worksheet.\n"))
            locators.append({"name": sheet.title, "rows": len(rows), "columns": width})
    finally:
        formulas.close()
        cached.close()
    if not any(s["columns"] for s in locators):
        raise ValueError("workbook has no populated cells")
    return "\n\n".join(parts), locators


def convert(root: Path, record):
    path = safe_path(root, root / record.path)
    warnings, inputs, metadata, locators = [], [], {}, {}
    title, converter = path.stem or path.name, f"docnorm.{record.kind}"
    if record.kind == "html":
        title, fragments, converter, metadata, locators = html_adapter(root, path, warnings)
    elif record.kind == "markdown":
        fragments = expand_markdown(root, path, warnings, inputs)
        title = title_for("".join(f.text for f in fragments), title)
    elif record.kind == "docx":
        check_office_archive(path)
        text = local_markitdown(path.read_bytes(), ".docx")
        fragments = [Fragment(text, path.parent)]
        title = title_for(text, title)
        converter = "MarkItDown.DocxConverter"
    elif record.kind == "xlsx":
        text, sheets = xlsx_adapter(path, warnings)
        fragments = [Fragment(text, path.parent)]
        locators = {"worksheets": sheets}
        converter = "docnorm.openpyxl-fidelity"
    else:
        text, ws = decode(path.read_bytes())
        warnings.extend(ws)
        if not text.strip():
            raise ValueError("empty source content")
        if record.kind == "json":
            obj = json.loads(text)
            if isinstance(obj, dict) and all(k in obj for k in ["schemaVersion", "entities", "components", "relations", "scenarios"]):
                fragments = [architecture_adapter(obj, path)]
                title = obj.get("title") or title
                metadata = {"schemaVersion": obj["schemaVersion"], "sourceRoot": obj.get("sourceRoot")}
                converter = "docnorm.architecture-json"
            else:
                fragments = [Fragment(fence(text, "json"), path.parent)]
        elif record.kind == "csv":
            try:
                dialect = csv.Sniffer().sniff(text[:8192])
            except csv.Error:
                dialect = csv.excel
            rows = list(csv.reader(io.StringIO(text), dialect, strict=True))
            if len({len(row) for row in rows}) > 1:
                warnings.append("CSV rows have inconsistent widths; blank padding retained")
            fragments = [Fragment(table(rows), path.parent)]
        elif record.kind == "text":
            fragments = [Fragment(text, path.parent)]
            title = title_for(text, title)
        else:
            fragments = [Fragment(fence(text, record.kind), path.parent)]
    combined = "".join(f.text for f in fragments)
    if not combined.strip() or not re.sub(r'<a id="[^"]*"></a>', "", combined).strip():
        raise ValueError("conversion produced empty content")
    return Converted(record.path, str(title), fragments, converter, warnings, metadata, locators, sorted(set(inputs)))


class AdapterRegistry:
    """Explicit local adapters only; there is no plugin discovery."""

    def __init__(self):
        self._adapters = {}

    def register(self, kinds, adapter):
        for kind in kinds:
            self._adapters[kind] = adapter

    def convert(self, root, record):
        adapter = self._adapters.get(record.kind)
        if adapter is None:
            raise ValueError(f"no adapter registered for {record.kind}")
        return adapter(root, record)


registry = AdapterRegistry()
registry.register(["markdown", "text", "html", "docx", "xlsx", "csv", "json", "yaml", "mermaid", "plantuml", "bash", "javascript"], convert)
